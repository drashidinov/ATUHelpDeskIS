import openpyxl, glob, json, re, sys
sys.path.insert(0, '/home/claude/parse2')
from rooms import find_target_rooms_in_text, ALL_ROOMS

UPLOAD_DIR = "/mnt/user-data/uploads"

FILES = [
    "2_курс_ФИиИТ_ФИиИС_Расписание_1_академического_периода_2026-2027уч_года.xlsx",
    "2_курс_ФПТ_ФБиХТ_Расписание_1_академического_периода_2026-2027_уч_г_.xlsx",
    "2курс_ФДТТиО_ФЭиБ_РАСПИСАНИЕ_2026-2027_уч_г.xlsx",
    "3_курс_2026-2027.xlsx",
    "4_курс_расписание_2026-2027_г_14_09.xlsx",
    "МАГИСТРАТУРА_1_КУРС_ПРОФИЛЬНОЕ_НАПР__2026-2027_г_.xlsx",
    "Расписание_докторантов_1_курс__1_.xlsx",
    "Расписание_магистрантов_1_КУРС_НАУЧНО-ПЕД__НАПРАВЛЕНИЕ__2026-2027_уч_г_.xlsx",
    "Расписание_магистрантов_2_курс.xlsx",
    "1г_о__ТИПО_ВО_2курс_ФДТТиО_ФЭиБ_РАСПИСАНИЕ_2026-2027_уч_г_ТИПО.xlsx",
    "ТИПО_3года__ВО_3года2_курс_ФПТ_ФБиХТ_ФИиИТ_ФИиИС_Расписание_1_академического_периода_2026-2027_уч_г_.xlsx",
    "14_09_Г_Ж_1_курс_2026-2027_ФПТ_ФБиХТ__1_.xlsx",
    "14_09_Г_Ж_1_курс_2026-2027_Фэб_фдтио_Фиит.xlsx",
]

def get_header_fill_map(ws, header_row_idx, max_col):
    fill = {}
    for c in range(1, max_col+1):
        v = ws.cell(row=header_row_idx, column=c).value
        fill[c] = v
    for mc in ws.merged_cells.ranges:
        if mc.min_row <= header_row_idx <= mc.max_row:
            top_val = ws.cell(row=mc.min_row, column=mc.min_col).value
            for c in range(mc.min_col, mc.max_col+1):
                if fill.get(c) is None:
                    fill[c] = top_val
    return fill

def clean_group_name(g):
    if not g:
        return None
    g = str(g).strip()
    g = re.sub(r'\s*\(\d+\)\s*$', '', g)
    return g.strip()

records = []

for fname in FILES:
    fpath = f"{UPLOAD_DIR}/{fname}"
    try:
        wb = openpyxl.load_workbook(fpath, data_only=True)
    except Exception as e:
        print("FAIL open", fname, e)
        continue
    for sn in wb.sheetnames:
        ws = wb[sn]
        max_row = ws.max_row
        max_col = ws.max_column
        if max_row is None or max_row < 3:
            continue
        header_row = None
        for r in range(1, min(15, max_row)+1):
            for c in range(1, min(4, max_col)+1):
                v = ws.cell(row=r, column=c).value
                if isinstance(v, str) and ('ремя' in v or 'ақыт' in v):
                    header_row = r
                    break
            if header_row:
                break
        if not header_row:
            continue

        header_map = get_header_fill_map(ws, header_row, max_col)

        current_day = None
        for r in range(header_row+1, max_row+1):
            day_cell = ws.cell(row=r, column=1).value
            if day_cell and str(day_cell).strip():
                current_day = str(day_cell).strip()
            time_cell = ws.cell(row=r, column=2).value
            time_str = str(time_cell).strip() if time_cell else None
            if not time_str or not current_day:
                continue
            for c in range(3, max_col+1):
                cell_val = ws.cell(row=r, column=c).value
                if not isinstance(cell_val, str) or not cell_val.strip():
                    continue
                rooms_found = find_target_rooms_in_text(cell_val)
                if not rooms_found:
                    continue
                group = clean_group_name(header_map.get(c))
                for room in rooms_found:
                    records.append({
                        "file": fname,
                        "sheet": sn,
                        "group": group,
                        "day": current_day,
                        "time": time_str,
                        "room": room,
                        "text": cell_val.strip(),
                    })
    wb.close()
    print("done", fname, "records so far:", len(records))

with open('/home/claude/parse2/records.json', 'w', encoding='utf-8') as f:
    json.dump(records, f, ensure_ascii=False, indent=1)

print("TOTAL RECORDS:", len(records))
